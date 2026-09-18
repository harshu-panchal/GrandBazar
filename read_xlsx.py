import openpyxl
import json

wb = openpyxl.load_workbook(r'd:\AppZeto\GrandBazar\Zinto_Test_Case_Scenarios_1.xlsx', read_only=True, data_only=True)

print("SHEETS:")
for i, s in enumerate(wb.sheetnames):
    print(f"{i}: {s}")

print("\n--- READING TEST CASE SHEETS ---")

test_sheets = [s for s in wb.sheetnames if s not in ['Summary', 'UAT Requirement Matrix', 'Go-Live Sign-off Checklist']]

all_tcs = []
for sheet_name in test_sheets:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    header_row_idx = 0
    headers = None
    for i, row in enumerate(rows[:5]):
        if row and any(str(c).strip().lower() in ['test case id', 'tc id', 'id'] for c in row if c):
            headers = [str(c).strip() if c else '' for c in row]
            header_row_idx = i
            break
    if not headers:
        headers = [str(c).strip() if c else '' for c in rows[0]]
        header_row_idx = 0
    
    print(f"\nSheet: {sheet_name} | Headers: {headers[:8]}")
    count = 0
    for row in rows[header_row_idx+1:]:
        if not any(row):
            continue
        row_data = dict(zip(headers, [str(c).strip() if c is not None else '' for c in row]))
        tc_id = row_data.get('Test Case ID', '') or row_data.get('TC ID', '') or row_data.get('ID', '')
        if tc_id and tc_id.upper() not in ('TEST CASE ID', 'TC ID', 'ID', ''):
            all_tcs.append({
                'sheet': sheet_name,
                'tc_id': tc_id,
                'module': row_data.get('Module', ''),
                'scenario': row_data.get('Test Scenario', '') or row_data.get('Scenario', ''),
                'priority': row_data.get('Priority', ''),
                'type': row_data.get('Type', '') or row_data.get('Test Type', ''),
                'actor': row_data.get('Actor / Role', '') or row_data.get('Actor', '') or row_data.get('Role', ''),
                'status': row_data.get('Status', ''),
            })
            count += 1
    print(f"  Test cases found: {count}")

print(f"\nTOTAL TEST CASES: {len(all_tcs)}")
print("\nFirst 10 TCs:")
for tc in all_tcs[:10]:
    print(tc)
