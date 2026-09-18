import openpyxl
import json

wb = openpyxl.load_workbook(r'd:\AppZeto\GrandBazar\Zinto_Test_Case_Scenarios_1.xlsx', read_only=True, data_only=True)

output_lines = []
output_lines.append("SHEETS:")
for i, s in enumerate(wb.sheetnames):
    output_lines.append(f"{i}: {s}")

test_sheets = [s for s in wb.sheetnames if s not in ['Summary', 'UAT Requirement Matrix', 'Go-Live Sign-off Checklist']]

all_tcs = []
full_tcs = []  # Full details

for sheet_name in test_sheets:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    header_row_idx = 0
    headers = None
    for i, row in enumerate(rows[:5]):
        if row and any(str(c).strip().lower() in ['test case id', 'tc id', 'id'] for c in row if c):
            headers = [str(c).strip() if c else '' for c in row]
            header_row_idx = i
            break
    if not headers:
        headers = [str(c).strip() if c else '' for c in rows[0]]
        header_row_idx = 0
    
    output_lines.append(f"\nSheet: {sheet_name}")
    output_lines.append(f"Headers: {headers}")
    count = 0
    for row in rows[header_row_idx+1:]:
        if not any(row):
            continue
        row_data = dict(zip(headers, [str(c).strip() if c is not None else '' for c in row]))
        tc_id = row_data.get('Test Case ID', '') or row_data.get('TC ID', '') or row_data.get('ID', '')
        if tc_id and tc_id.upper() not in ('TEST CASE ID', 'TC ID', 'ID', '', 'NONE'):
            tc_info = {
                'sheet': sheet_name,
                'tc_id': tc_id,
                'module': row_data.get('Module', ''),
                'sub_feature': row_data.get('Sub-Feature / Flow', '') or row_data.get('Sub-Feature', ''),
                'scenario': row_data.get('Test Scenario', '') or row_data.get('Scenario', ''),
                'preconditions': row_data.get('Preconditions', ''),
                'steps': row_data.get('Test Steps', '') or row_data.get('Steps', ''),
                'test_data': row_data.get('Test Data', ''),
                'expected': row_data.get('Expected Result', '') or row_data.get('Expected', ''),
                'priority': row_data.get('Priority', ''),
                'type': row_data.get('Type', '') or row_data.get('Test Type', ''),
                'actor': row_data.get('Actor / Role', '') or row_data.get('Actor', '') or row_data.get('Role', ''),
                'status': row_data.get('Status', ''),
            }
            all_tcs.append(tc_info)
            count += 1
    output_lines.append(f"Test cases found: {count}")

output_lines.append(f"\nTOTAL TEST CASES: {len(all_tcs)}")

# By priority
from collections import Counter
priorities = Counter(tc['priority'] for tc in all_tcs)
output_lines.append(f"\nBy Priority: {dict(priorities)}")

# By sheet
by_sheet = Counter(tc['sheet'] for tc in all_tcs)
output_lines.append(f"\nBy Sheet: {dict(by_sheet)}")

# By actor
by_actor = Counter(tc['actor'] for tc in all_tcs)
output_lines.append(f"\nBy Actor: {dict(by_actor)}")

# By status
by_status = Counter(tc['status'] for tc in all_tcs)
output_lines.append(f"\nBy Status (existing): {dict(by_status)}")

# Save all test cases to JSON
with open(r'd:\AppZeto\GrandBazar\all_test_cases.json', 'w', encoding='utf-8') as f:
    json.dump(all_tcs, f, indent=2, ensure_ascii=False)

with open(r'd:\AppZeto\GrandBazar\xlsx_analysis.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(output_lines))

print("Done. Files written.")
print('\n'.join(output_lines[:100]))
